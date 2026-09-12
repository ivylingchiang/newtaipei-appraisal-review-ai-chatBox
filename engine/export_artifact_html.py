# -*- coding: utf-8 -*-
"""thirdVersion 表單預覽 Artifact：渲染實際填好的 xlsx，逐格可點看填表依據。"""
import json, html, os, collections
import openpyxl

ROOT = "/Users/ivychiang/Code/NewTaipei/newtaipei-appraisal-review-ai-chatbox"
SRC = os.path.join(ROOT, "output/thirdVersion")
INF = json.load(open(os.path.join(ROOT, "datasets/external/poi_inference.json"), encoding="utf-8"))
E = lambda s: html.escape(str(s) if s is not None else "")

# Excel 底色 → (代號, 顯示名, 說明)
KIND = {
    "D9EAD3": ("orig", "題目原載", "doc/題目.pdf 原已載明之勘查事實或價格，逕予抄錄"),
    "FFF2CC": ("ai", "AI判定", "題目事實 × 評價基準明細表查表，無估算成分"),
    "FCE5CD": ("pres", "推定", "以區段層級記載推定至宗地層級，未經個別勘查"),
    "CFE2F3": ("ext", "外部推算 · 官方", "開放資料推算，官方名冊來源（信心 A／B）"),
    "E1D5E7": ("osm", "外部推算 · OSM", "僅 OpenStreetMap 群眾協作（信心 C）"),
    "F9CB9C": ("blk", "覆蓋不足", "嫌惡設施子欄缺漏，推算等級僅為樂觀上限，不得判級"),
    "F4CCCC": ("gap", "資料不足", "應填而無可用資料，留白並列為退補事項"),
    "EFEFEF": ("na", "不適用", "普通住宅用地之評價基準明細表無此修正細項"),
}

SHEETS = [
    ("表3_地價區段勘查表_已填.xlsx", "表3 地價區段勘查表", True),
    ("表5-1_影響地價區域因素分析明細表_已填.xlsx", "表5-1 區域因素分析明細表", False),
    ("表4_比較法調查估價表_已填.xlsx", "表4 比較法調查估價表", False),
]


def spans(ws):
    m = {}
    for rg in ws.merged_cells.ranges:
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                m[(r, c)] = ("skip" if (r, c) != (rg.min_row, rg.min_col)
                             else (rg.max_row - rg.min_row + 1, rg.max_col - rg.min_col + 1))
    return m


def render(ws, notes):
    sm = spans(ws)
    counts = collections.Counter()
    out = ['<table class="xl">']
    for r in range(1, ws.max_row + 1):
        out.append("<tr>")
        for c in range(1, ws.max_column + 1):
            info = sm.get((r, c))
            if info == "skip":
                continue
            cell = ws.cell(r, c)
            a = []
            if info:
                rs, cs = info
                if rs > 1: a.append(f'rowspan="{rs}"')
                if cs > 1: a.append(f'colspan="{cs}"')
            rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
            klass = []
            if isinstance(rgb, str) and len(rgb) == 8 and rgb[2:] in KIND:
                k = KIND[rgb[2:]][0]
                klass.append("k-" + k)
                counts[k] += 1
            if cell.comment:
                nid = len(notes)
                notes.append(cell.comment.text)
                a.append(f'data-n="{nid}"')
                klass.append("hn")
            if klass:
                a.append(f'class="{" ".join(klass)}"')
            a.append(f'data-ref="{cell.coordinate}"')
            v = cell.value
            if v is None:
                v = ""
            elif isinstance(v, float):
                v = f"{v:,.2f}".rstrip("0").rstrip(".") if v % 1 else f"{v:,.0f}"
            out.append(f'<td {" ".join(a)}>{E(v).replace(chr(10), "<br>")}</td>')
        out.append("</tr>")
    out.append("</table>")
    return "\n".join(out), counts


notes, panels, tabs = [], [], []
first = True
for fn, title, per_seg in SHEETS:
    wb = openpyxl.load_workbook(os.path.join(SRC, fn))
    for ws in wb.worksheets:
        if ws.title == "填表依據":
            continue
        tid = f"t{len(tabs)}"
        label = ws.title if per_seg else title.split()[0]
        sub = title if not per_seg else "表3"
        tbl, cnt = render(ws, notes)
        chips = "".join(
            f'<span class="chip k-{k}">{E(KIND[[h for h in KIND if KIND[h][0]==k][0]][1])}'
            f'<b>{n}</b></span>'
            for k, n in sorted(cnt.items(), key=lambda kv: -kv[1]))
        tabs.append(f'<button class="tab{" on" if first else ""}" data-t="{tid}" '
                    f'role="tab" aria-selected="{"true" if first else "false"}" '
                    f'aria-controls="{tid}" id="tb-{tid}">'
                    f'<b>{E(label)}</b><i>{E(sub)}</i></button>')
        panels.append(f'<section class="pane{" on" if first else ""}" id="{tid}" '
                      f'role="tabpanel" aria-labelledby="tb-{tid}"{"" if first else " hidden"}>'
                      f'<div class="chips">{chips}</div>'
                      f'<div class="scroll">{tbl}</div></section>')
        first = False

NOTES_JSON = json.dumps(notes, ensure_ascii=False)

legend = "".join(
    f'<span class="lg k-{k}"><i></i><b>{E(nm)}</b><em>{E(desc)}</em></span>'
    for hexv, (k, nm, desc) in KIND.items())

segstrip = "".join(
    f'<div class="sg{" bench" if s=="P001-00" else ""}"><b>{s}</b>'
    f'<span class="xy">{INF["segments"][s]["center_twd97"][0]:,.0f} · '
    f'{INF["segments"][s]["center_twd97"][1]:,.0f}</span>'
    f'<span class="un">±{INF["segments"][s]["uncertainty_m"]:g} m</span></div>'
    for s in ["P001-00", "P002-00", "P003-00", "P004-00"])

CSS = """
:root{
 --ground:#F5F7F3;--surface:#FFFFFF;--sunk:#EAEEE7;
 --ink:#19211C;--ink2:#3C4842;--muted:#6B776F;--faint:#93A098;
 --rule:#D5DCD2;--rule2:#BFC9BB;--accent:#0E5C4A;--accent-w:#E1EDE8;
 --k-orig:#D9EAD3;--k-ai:#FFF2CC;--k-pres:#FCE5CD;--k-ext:#CFE2F3;
 --k-osm:#E1D5E7;--k-blk:#F9CB9C;--k-gap:#F4CCCC;--k-na:#EFEFEF;
 --cell-ink:#19211C;--blk:#B26A1F;
 --shadow:0 1px 2px rgba(25,33,28,.07),0 6px 20px rgba(25,33,28,.06);
}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
 --ground:#111614;--surface:#191F1B;--sunk:#222A25;
 --ink:#E8EDE7;--ink2:#C2CCC4;--muted:#93A099;--faint:#6C7872;
 --rule:#2C352E;--rule2:#3C4740;--accent:#63BFA3;--accent-w:#1D2C26;
 --k-orig:#2E4A2C;--k-ai:#4D4422;--k-pres:#4E3A25;--k-ext:#24445C;
 --k-osm:#3E3352;--k-blk:#59391A;--k-gap:#542B28;--k-na:#2A322C;
 --cell-ink:#EAEFE9;--blk:#E0A05A;
 --shadow:0 1px 2px rgba(0,0,0,.5),0 6px 20px rgba(0,0,0,.35);
}}
:root[data-theme="dark"]{
 --ground:#111614;--surface:#191F1B;--sunk:#222A25;
 --ink:#E8EDE7;--ink2:#C2CCC4;--muted:#93A099;--faint:#6C7872;
 --rule:#2C352E;--rule2:#3C4740;--accent:#63BFA3;--accent-w:#1D2C26;
 --k-orig:#2E4A2C;--k-ai:#4D4422;--k-pres:#4E3A25;--k-ext:#24445C;
 --k-osm:#3E3352;--k-blk:#59391A;--k-gap:#542B28;--k-na:#2A322C;
 --cell-ink:#EAEFE9;--blk:#E0A05A;
 --shadow:0 1px 2px rgba(0,0,0,.5),0 6px 20px rgba(0,0,0,.35);
}
*{box-sizing:border-box}
body{background:var(--ground);color:var(--ink);margin:0;padding-block:0 0;
 font:15px/1.7 "Noto Sans TC",-apple-system,"PingFang TC","Microsoft JhengHei",sans-serif;
 -webkit-font-smoothing:antialiased}
.wrap{max-width:1320px;margin:0 auto;padding-inline:18px}
h1,h2{font-family:"Noto Serif TC",serif;margin:0;text-wrap:balance}
.mono,code,.xy,.un,.chip b,.lg b{font-family:"IBM Plex Mono",ui-monospace,monospace;
 font-variant-numeric:tabular-nums}

.mast{padding-block:30px 0}
.eyebrow{font-size:11px;letter-spacing:.22em;color:var(--accent);font-weight:700;margin-bottom:9px}
.mast h1{font-size:clamp(23px,3.6vw,32px);line-height:1.3;font-weight:700}
.mast p{color:var(--ink2);max-width:66ch;margin:11px 0 0;font-size:14.5px}
.warn{margin-top:18px;background:var(--surface);border:1px solid var(--rule2);
 border-left:4px solid var(--blk);padding:13px 17px;font-size:13.5px;color:var(--ink2);
 line-height:1.7;max-width:none}
.warn b{color:var(--ink)}

.segs{display:flex;gap:8px;flex-wrap:wrap;margin-top:18px}
.sg{background:var(--surface);border:1px solid var(--rule);padding:7px 13px;
 display:flex;gap:11px;align-items:baseline;font-size:12px}
.sg.bench{border-color:var(--accent);box-shadow:inset 3px 0 0 var(--accent)}
.sg b{font-family:"IBM Plex Mono",monospace;font-size:12.5px}
.sg .xy{color:var(--muted);font-size:11.5px}
.sg .un{color:var(--accent);font-weight:600;font-size:11.5px}

.legend{display:grid;gap:6px;grid-template-columns:repeat(auto-fit,minmax(226px,1fr));
 margin-top:16px}
.lg{display:flex;align-items:baseline;gap:8px;font-size:12px;background:var(--surface);
 border:1px solid var(--rule);padding:6px 10px;line-height:1.5}
.lg i{width:12px;height:12px;flex:none;border:1px solid var(--rule2);display:block;
 position:relative;top:2px}
.lg b{font-size:11.5px;white-space:nowrap;font-weight:600}
.lg em{font-style:normal;color:var(--faint);font-size:11px}
.lg.k-orig i{background:var(--k-orig)} .lg.k-ai i{background:var(--k-ai)}
.lg.k-pres i{background:var(--k-pres)} .lg.k-ext i{background:var(--k-ext)}
.lg.k-osm i{background:var(--k-osm)} .lg.k-blk i{background:var(--k-blk)}
.lg.k-gap i{background:var(--k-gap)} .lg.k-na i{background:var(--k-na)}

.tabs{display:flex;gap:0;margin-top:26px;border-bottom:2px solid var(--rule2);
 overflow-x:auto;scrollbar-width:thin}
.tab{background:none;border:0;border-bottom:2px solid transparent;margin-bottom:-2px;
 padding:9px 15px;cursor:pointer;color:var(--muted);font-family:inherit;
 display:flex;flex-direction:column;align-items:flex-start;gap:0;white-space:nowrap;
 line-height:1.35}
.tab b{font-family:"IBM Plex Mono",monospace;font-size:13.5px;font-weight:600}
.tab i{font-style:normal;font-size:10.5px;letter-spacing:.05em}
.tab:hover{color:var(--ink)}
.tab.on{color:var(--accent);border-bottom-color:var(--accent)}
.tab:focus-visible{outline:2px solid var(--accent);outline-offset:-2px}

.pane{margin-top:14px}
.pane[hidden]{display:none}
.chips{display:flex;flex-wrap:wrap;gap:6px;margin-bottom:10px}
.chip{font-size:11px;padding:2px 9px;border:1px solid var(--rule2);display:inline-flex;
 gap:6px;align-items:baseline;color:var(--ink2)}
.chip b{font-size:11.5px}
.chip.k-orig{background:var(--k-orig)} .chip.k-ai{background:var(--k-ai)}
.chip.k-pres{background:var(--k-pres)} .chip.k-ext{background:var(--k-ext)}
.chip.k-osm{background:var(--k-osm)} .chip.k-blk{background:var(--k-blk)}
.chip.k-gap{background:var(--k-gap)} .chip.k-na{background:var(--k-na)}

.scroll{overflow-x:auto;background:var(--surface);border:1px solid var(--rule2);
 box-shadow:var(--shadow)}
table.xl{border-collapse:collapse;font-size:11px;white-space:nowrap;color:var(--cell-ink)}
table.xl td{border:1px solid var(--rule);padding:3px 6px;vertical-align:middle;
 max-width:330px;overflow:hidden;text-overflow:ellipsis;line-height:1.5}
td.k-orig{background:var(--k-orig)} td.k-ai{background:var(--k-ai)}
td.k-pres{background:var(--k-pres)} td.k-ext{background:var(--k-ext)}
td.k-osm{background:var(--k-osm)} td.k-blk{background:var(--k-blk)}
td.k-gap{background:var(--k-gap)} td.k-na{background:var(--k-na)}
td.hn{cursor:pointer;position:relative}
td.hn::after{content:"";position:absolute;top:1px;right:1px;width:0;height:0;
 border-top:5px solid var(--accent);border-left:5px solid transparent}
td.sel{outline:2px solid var(--accent);outline-offset:-2px}

.panel{position:sticky;bottom:0;background:var(--surface);border-top:2px solid var(--accent);
 box-shadow:0 -4px 18px rgba(0,0,0,.1);padding:13px 18px;margin-top:0;z-index:5}
.panel .ph{display:flex;gap:11px;align-items:baseline;margin-bottom:6px}
.panel .ref{font-family:"IBM Plex Mono",monospace;font-size:12px;background:var(--accent);
 color:var(--ground);padding:1px 8px;font-weight:600}
.panel .hint{color:var(--faint);font-size:12px}
.panel pre{margin:0;white-space:pre-wrap;font-size:12.5px;line-height:1.75;
 font-family:inherit;color:var(--ink2);max-height:148px;overflow-y:auto}
.panel button.x{margin-left:auto;background:none;border:1px solid var(--rule2);
 color:var(--muted);cursor:pointer;font-size:11px;padding:2px 9px;font-family:inherit}

footer{margin-top:34px;padding-block:18px 40px;border-top:1px solid var(--rule2);
 font-size:11.5px;color:var(--faint);line-height:1.9}
footer code{font-size:11px}
@media (max-width:640px){
 .legend{grid-template-columns:1fr} .lg em{display:none}
 .panel pre{max-height:110px}
}
@media (prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
"""

JS = """
const NOTES=%s;
const tabs=[...document.querySelectorAll('.tab')];
tabs.forEach(t=>t.addEventListener('click',()=>{
 tabs.forEach(o=>{o.classList.toggle('on',o===t);o.setAttribute('aria-selected',o===t);});
 document.querySelectorAll('.pane').forEach(p=>{
  const on=p.id===t.dataset.t; p.classList.toggle('on',on); p.hidden=!on;});
}));
const panel=document.getElementById('panel'),pref=document.getElementById('pref'),
      ptxt=document.getElementById('ptxt');
let sel=null;
document.addEventListener('click',ev=>{
 const td=ev.target.closest('td.hn'); if(!td)return;
 if(sel)sel.classList.remove('sel'); sel=td; td.classList.add('sel');
 pref.textContent=td.dataset.ref; ptxt.textContent=NOTES[+td.dataset.n];
 panel.hidden=false;
});
document.getElementById('pclose').addEventListener('click',()=>{
 panel.hidden=true; if(sel){sel.classList.remove('sel');sel=null;}
});
""" % NOTES_JSON

DOC = f"""<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;600;700&family=Noto+Serif+TC:wght@600;700&family=IBM+Plex+Mono:wght@400;600&display=swap">
<title>樹林區表3 填表預覽</title>
<style>{CSS}</style>
<div class="wrap">
<div class="mast">
 <div class="eyebrow">OUTPUT / THIRDVERSION　·　新北市樹林區　普通住宅用地　案號 1110901-99-XXX</div>
 <h1>表3／表5／表4 填表預覽</h1>
 <p>四份表3 的設施類欄位原本全部空白，本版以 19 個開放資料來源推算補填。
 <b>有底色的格子點一下，下方會顯示該格的填表依據</b>（與 Excel 內的儲存格註解相同）。</p>
 <div class="warn"><b>推算值不是勘查記錄。</b>表3 是法定勘查記錄，權威性來自承辦單位實地認定；藍／紫格是第三方開放資料推算值，
 距離為區段中心點<b>直線距離</b>（正向設施因此系統性高估便利性）。可作外業候選清單、審查合理性檢核與退補依據，
 <b>不可直接送件，更不可據以認定某設施「無」</b>。紅格代表查無資料而留白 —— 沒有任何一格被填成「無」。</div>
 <div class="segs">{segstrip}</div>
 <div class="legend">{legend}</div>
</div>

<div class="tabs" role="tablist">{''.join(tabs)}</div>
{''.join(panels)}

<footer>
 區段中心點由 OpenStreetMap 界街幾何交會求得，不確定半徑 25–66 m，小於所有最內門檻（200–300 m）。
 P001-00 推算中心距捷運 LG16 站 72 m，與其 <code>scope_desc</code> 所載「捷運開發區」自洽。<br>
 級距與修正率：新北市樹林區普通住宅用地影響地價區域因素評價基準明細表　·　填表規範：內政部土地徵收補償市價查估作業手冊（104年3月）<br>
 產出：<code>engine/poi_fetch.py</code> → <code>poi_infer.py</code> → <code>export_v3.py</code>　·　逐格明細見各 xlsx 的「填表依據」工作表　·　推算日期 {INF['generated_at']}
</footer>
</div>

<div class="panel" id="panel" hidden>
 <div class="ph"><span class="ref" id="pref">—</span>
  <span class="hint">填表依據</span>
  <button class="x" id="pclose">關閉</button></div>
 <pre id="ptxt"></pre>
</div>
<script>{JS}</script>
"""
out = os.environ.get("OUT_HTML") or os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "v3.html")
open(out, "w", encoding="utf-8").write(DOC)
print("→", out, f"{len(DOC)/1024:.0f} KB", "| notes:", len(notes), "| panes:", len(panels))
