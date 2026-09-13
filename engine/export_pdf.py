#!/usr/bin/env python3
"""把已填的 xlsx 排成黑白 PDF——本機沒有 Excel／LibreOffice 時的輸出路徑。

排版依工作表自身的版面資料：欄寬、列高、合併儲存格、框線樣式、對齊方式、
直書（textRotation 255）與數值格式，並比照 Excel 的「調整為一頁」縮放。
只畫框線與文字，不畫底色，與官方書表範本的外觀一致。

    python3 engine/export_pdf.py <xlsx> [<pdf>]
"""
import os
import re
import sys

import matplotlib
matplotlib.use("pdf")
# Type 42：嵌入 TrueType 子集，PDF 內的文字才能被搜尋與複製（預設的 Type 3 不行）
matplotlib.rcParams["pdf.fonttype"] = 42
from matplotlib import pyplot as plt                                # noqa: E402
from matplotlib import font_manager                                 # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages                # noqa: E402
from matplotlib.patheffects import withStroke                       # noqa: E402
from PIL import ImageFont                                           # noqa: E402
import openpyxl                                                     # noqa: E402
from openpyxl.utils import get_column_letter                        # noqa: E402

# 內建 CJK 字型：macOS 全系統皆有，涵蓋表單用到的 ●○■□─ 等符號
FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Songti.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
]
FONT_PATH = next((p for p in FONT_CANDIDATES if os.path.exists(p)), None)
if FONT_PATH is None:
    raise SystemExit("找不到可用的中文字型")
FP = font_manager.FontProperties(fname=FONT_PATH)
_MEASURE = ImageFont.truetype(FONT_PATH, 100)                # 量測用，寬度與字級成正比

A4 = (595.276, 841.890)          # points
PT_PER_IN = 72.0
DEFAULT_COL_W = 8.43             # Excel 欄寬單位（字元）
DEFAULT_ROW_H = 15.0             # points
MIN_FONT = 4.0
PAD = 1.6                        # 儲存格內距（pt，未縮放）

LINE_W = {"hair": 0.25, "thin": 0.5, "medium": 1.1, "thick": 1.7,
          "double": 1.1, "dotted": 0.5, "dashed": 0.5, "dashDot": 0.5,
          "dashDotDot": 0.5, "mediumDashed": 1.1, "slantDashDot": 0.5,
          "mediumDashDot": 1.1, "mediumDashDotDot": 1.1}
DASHED = {"dotted": (0.5, (0.5, 1.2)), "dashed": (0.5, (2.5, 1.5)),
          "dashDot": (0.5, (3, 1.2, 0.8, 1.2)),
          "dashDotDot": (0.5, (3, 1.2, 0.8, 1.2, 0.8, 1.2)),
          "mediumDashed": (1.1, (2.5, 1.5))}

CJK = re.compile(r"[ᄀ-ᇿ⺀-꓏가-힣豈-﫿"
                 r"︰-﹏＀-￯　-〿]")
TOKEN = re.compile(r"[0-9][0-9,.]*%?|[A-Za-z]+|\s+|.", re.S)


# ------------------------------------------------------------------ 版面度量
def col_width_pt(ws, idx):
    d = ws.column_dimensions.get(get_column_letter(idx))
    w = d.width if d is not None and d.width else (
        ws.sheet_format.defaultColWidth or DEFAULT_COL_W)
    return (round(w * 7) + 5) * 0.75          # 字元 → px → pt


def row_height_pt(ws, idx):
    d = ws.row_dimensions.get(idx)
    h = d.height if d is not None and d.height else (
        ws.sheet_format.defaultRowHeight or DEFAULT_ROW_H)
    return float(h)


def used_bounds(ws):
    """列印範圍；未設定時取「有內容」的範圍，空白欄列不算（框線會延伸到表外）"""
    pa = ws.print_area
    if pa:
        rng = openpyxl.worksheet.cell_range.CellRange(
            pa[0] if isinstance(pa, (list, tuple)) else pa.split("!")[-1])
        return rng.max_row, rng.max_col
    max_r = max_c = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                max_r, max_c = max(max_r, c.row), max(max_c, c.column)
    for m in ws.merged_cells.ranges:                 # 有值之合併範圍須完整涵蓋
        if ws.cell(m.min_row, m.min_col).value not in (None, ""):
            max_r, max_c = max(max_r, m.max_row), max(max_c, m.max_col)
    return max(max_r, 1), max(max_c, 1)


def span_map(ws):
    """(row, col) → ('anchor', rows, cols) 或 'skip'"""
    m = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = ("skip" if (r, c) != (rng.min_row, rng.min_col)
                             else ("anchor", rng.max_row - rng.min_row + 1,
                                   rng.max_col - rng.min_col + 1))
    return m


def edge_map(ws):
    """(row, col) → 該格在合併範圍內不該畫的邊；合併範圍的內部框線 Excel 不畫"""
    m = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = {
                    "top": r != rng.min_row, "bottom": r != rng.max_row,
                    "left": c != rng.min_col, "right": c != rng.max_col}
    return m


# ------------------------------------------------------------------ 值的呈現
def fmt_number(v, fmt):
    fmt = (fmt or "General").split(";")[0]
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    if "%" in fmt:
        dec = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
        return f"{v * 100:.{dec}f}%"
    if fmt in ("General", "@"):
        return f"{v:g}" if isinstance(v, float) else str(v)
    comma = "#,##" in fmt or "," in fmt.replace("#,##", "")
    dec = 0
    if "." in fmt:
        dec = len(re.match(r"[0#]*", fmt.split(".")[1]).group(0))
    s = f"{v:,.{dec}f}" if comma else f"{v:.{dec}f}"
    return s


def cell_text(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return fmt_number(v, cell.number_format)


# ------------------------------------------------------------------ 文字排版
def text_width(s, size):
    return _MEASURE.getlength(s) * size / 100.0


def wrap(text, size, avail):
    """回傳折行後的每一行。中文逐字可斷，英數以字串為單位。"""
    lines = []
    for para in str(text).split("\n"):
        if not para:
            lines.append("")
            continue
        cur = ""
        for tok in TOKEN.findall(para):
            if tok == "\t":
                tok = "    "
            trial = cur + tok
            if cur and text_width(trial, size) > avail:
                lines.append(cur.rstrip() if not CJK.search(cur[-1:]) else cur)
                cur = "" if tok.isspace() else tok
            else:
                cur = trial
        lines.append(cur)
    return lines or [""]


def fit_lines(text, size, avail_w, avail_h, wrap_text):
    """縮字級直到塞得下；回傳 (每行文字, 實際字級)"""
    s = size
    while s > MIN_FONT:
        lines = (wrap(text, s, avail_w) if wrap_text
                 else [ln for ln in str(text).split("\n")])
        too_wide = (not wrap_text
                    and max(text_width(ln, s) for ln in lines) > avail_w)
        if len(lines) * s * 1.18 <= avail_h and not too_wide:
            return lines, s
        s -= 0.5 if s > 6 else 0.25
    return (wrap(text, s, avail_w) if wrap_text else str(text).split("\n")), s


# ------------------------------------------------------------------ 繪製
def draw_borders(ax, ws, em, xs, ys, max_r, max_c, k):
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(r, c)
            b = cell.border
            inner = em.get((r, c))
            x0, x1 = xs[c - 1], xs[c]
            y0, y1 = ys[r - 1], ys[r]
            for side, (p0, p1) in (("left", ((x0, y0), (x0, y1))),
                                   ("right", ((x1, y0), (x1, y1))),
                                   ("top", ((x0, y0), (x1, y0))),
                                   ("bottom", ((x0, y1), (x1, y1)))):
                st = getattr(b, side).style
                if not st or (inner and inner[side]):
                    continue
                lw = LINE_W.get(st, 0.5) * k
                kw = {}
                if st in DASHED:
                    kw["dashes"] = tuple(d * k for d in DASHED[st][1])
                ax.plot([p0[0], p1[0]], [p0[1], p1[1]], color="black",
                        linewidth=lw, solid_capstyle="projecting", **kw)
                if st == "double":
                    off = 1.2 * k
                    dx, dy = (off, 0) if side in ("left", "right") else (0, off)
                    for sgn in (-1, 1):
                        ax.plot([p0[0] + sgn * dx, p1[0] + sgn * dx],
                                [p0[1] + sgn * dy, p1[1] + sgn * dy],
                                color="black", linewidth=0.35 * k)


def blank(ws, sm, r, c):
    """該格在畫面上是否為空白（合併範圍以其左上角的值為準）"""
    info = sm.get((r, c))
    if info == "skip":
        for m in ws.merged_cells.ranges:
            if (m.min_row <= r <= m.max_row) and (m.min_col <= c <= m.max_col):
                return ws.cell(m.min_row, m.min_col).value in (None, "")
    return ws.cell(r, c).value in (None, "")


def overflow_span(ws, sm, r, c, cs, max_c, ha, need, xs, pad):
    """Excel 的溢出顯示：文字過長時往相鄰空白格延伸，回傳 (x0, x1)"""
    lo, hi = c, c + cs - 1
    grow_l = ha in ("right", "center")
    grow_r = ha in ("left", "center", "general", None)
    while (xs[hi] - xs[lo - 1]) - 2 * pad < need:
        moved = False
        if grow_r and hi < max_c and blank(ws, sm, r, hi + 1):
            hi += 1
            moved = True
        if grow_l and lo > 1 and blank(ws, sm, r, lo - 1):
            lo -= 1
            moved = True
        if not moved:
            break
    return xs[lo - 1], xs[hi]


def draw_text(ax, ws, sm, xs, ys, max_r, max_c, k):
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            info = sm.get((r, c))
            if info == "skip":
                continue
            cell = ws.cell(r, c)
            txt = cell_text(cell)
            if not txt.strip():
                continue
            rs, cs = (info[1], info[2]) if info else (1, 1)
            x0, x1 = xs[c - 1], xs[min(c - 1 + cs, max_c)]
            y0, y1 = ys[r - 1], ys[min(r - 1 + rs, max_r)]
            al = cell.alignment
            size = (cell.font.sz or 11) * k
            pad = PAD * k
            w, h = (x1 - x0) - 2 * pad, (y1 - y0) - 2 * pad

            if al.textRotation == 255:              # 直書：逐字往下排
                chars = [ch for ch in txt if ch.strip()]
                s = min(size, max(MIN_FONT, h / max(len(chars), 1) / 1.12))
                total = len(chars) * s * 1.12
                cy = (y0 + y1) / 2 - total / 2 + s * 0.56
                for ch in chars:
                    ax.text((x0 + x1) / 2, cy, ch, ha="center", va="center",
                            fontproperties=FP, fontsize=s, color="black")
                    cy += s * 1.12
                continue

            ha = al.horizontal
            if ha in (None, "general"):
                ha = "right" if isinstance(cell.value, (int, float)) and not \
                    isinstance(cell.value, bool) else "left"
            if ha in ("centerContinuous", "distributed", "justify"):
                ha = "center"

            wrap_text = bool(al.wrap_text) or "\n" in txt
            numeric = isinstance(cell.value, (int, float)) and not isinstance(
                cell.value, bool)
            if not wrap_text and not numeric and text_width(txt, size) > w:
                # 未設定自動換行者比照 Excel 溢出到相鄰空白格；仍不夠則改為換行
                x0, x1 = overflow_span(ws, sm, r, c, cs, max_c, ha,
                                       text_width(txt, size), xs, pad)
                w = (x1 - x0) - 2 * pad
                wrap_text = text_width(txt, size) > w * 1.18
            lines, s = fit_lines(txt, size, w, h, wrap_text)
            lh = s * 1.18
            lines = lines[:max(1, int((h + lh * 0.35) // lh))]   # 不畫出格外
            # 表單格多為單行，未指定時置中比 Excel 的預設靠下更貼近印製稿
            va = al.vertical or "center"

            block = len(lines) * lh
            if va == "top":
                ty = y0 + pad + lh * 0.78
            elif va in ("center", "distributed", "justify"):
                ty = (y0 + y1) / 2 - block / 2 + lh * 0.78
            else:
                ty = y1 - pad - block + lh * 0.78
            tx = {"left": x0 + pad, "right": x1 - pad,
                  "center": (x0 + x1) / 2}[ha]
            mha = {"left": "left", "right": "right", "center": "center"}[ha]
            fx = [withStroke(linewidth=s * 0.045, foreground="black")] \
                if cell.font.b else None
            for ln in lines:
                if ln.strip():
                    t = ax.text(tx, ty, ln, ha=mha, va="baseline",
                                fontproperties=FP, fontsize=s, color="black",
                                path_effects=fx)
                    t.set_clip_path(plt.Rectangle(
                        (x0, y0), x1 - x0, y1 - y0,
                        transform=ax.transData, figure=ax.figure))
                ty += lh


def render_sheet(pdf, ws):
    max_r, max_c = used_bounds(ws)
    wcols = [col_width_pt(ws, c) for c in range(1, max_c + 1)]
    hrows = [row_height_pt(ws, r) for r in range(1, max_r + 1)]
    total_w, total_h = sum(wcols), sum(hrows)

    landscape = (ws.page_setup.orientation == "landscape")
    pw, ph = (A4[1], A4[0]) if landscape else A4
    pm = ws.page_margins
    ml, mr = (pm.left or 0.25) * PT_PER_IN, (pm.right or 0.25) * PT_PER_IN
    mt, mb = (pm.top or 0.4) * PT_PER_IN, (pm.bottom or 0.4) * PT_PER_IN
    k = min((pw - ml - mr) / total_w, (ph - mt - mb) / total_h)

    xs = [ml]
    for w in wcols:
        xs.append(xs[-1] + w * k)
    ys = [mt]
    for h in hrows:
        ys.append(ys[-1] + h * k)

    fig = plt.figure(figsize=(pw / PT_PER_IN, ph / PT_PER_IN))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, pw)
    ax.set_ylim(ph, 0)
    ax.axis("off")
    sm, em = span_map(ws), edge_map(ws)
    draw_borders(ax, ws, em, xs, ys, max_r, max_c, k)
    draw_text(ax, ws, sm, xs, ys, max_r, max_c, k)
    pdf.savefig(fig)
    plt.close(fig)


def render(xlsx_path, pdf_path, skip_sheets=()):
    wb = openpyxl.load_workbook(xlsx_path)
    sheets = [ws for ws in wb.worksheets
              if ws.title not in skip_sheets and ws.sheet_state == "visible"]
    with PdfPages(pdf_path) as pdf:
        for ws in sheets:
            render_sheet(pdf, ws)
        pdf.infodict()["Title"] = os.path.splitext(os.path.basename(pdf_path))[0]
    return len(sheets)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".pdf"
    n = render(src, dst, skip_sheets=("填表依據",))
    print(f"{dst}（{n} 頁）")
