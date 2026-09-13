#!/usr/bin/env python3
"""output/thirdVersion — 以開放資料推算補填表3 設施類欄位後之表3／表5。

與 secondVersion 的唯一差別：表3 的 13 個設施類細項不再留白，改以
datasets/external/poi_inference.json 之推算結果填入，並逐格標示信心等級。

⚠️ 推算值屬【外部資料佐證】，不是法定勘查記錄，不得取代承辦單位實地勘查。
   理由與限制見 output/thirdVersion/README.md §0。
"""
import os, sys, json
import openpyxl
from openpyxl.styles import PatternFill, Alignment
from openpyxl.comments import Comment

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import export_xlsx as X
from loader import Dataset
from compute import derive_segment_levels, derive_table4_individual

ROOT = X.ROOT
OUT = os.path.join(ROOT, "output", "thirdVersion")
INF = os.path.join(ROOT, "datasets", "external", "poi_inference.json")
REGION = X.REGION
os.makedirs(OUT, exist_ok=True)
X.OUT = OUT

# ------------------------------------------------ 新增底色：外部資料推算三級
X.FILL["外部推算"] = PatternFill("solid", fgColor="CFE2F3")      # 藍：官方名冊(A/B)
X.FILL["外部推算OSM"] = PatternFill("solid", fgColor="E1D5E7")   # 紫：僅 OSM 群眾協作(C)
X.FILL["覆蓋不足"] = PatternFill("solid", fgColor="F9CB9C")      # 橘：子欄缺漏，判級不可用
X.LEGEND = X.LEGEND + [
    ("外部推算", "CFE2F3", "開放資料推算，官方名冊來源（信心 A／B）。距離為區段中心點至最近"
                          "設施之直線距離，非現場實測，僅供核對"),
    ("外部推算OSM", "E1D5E7", "開放資料推算，僅 OpenStreetMap 群眾協作來源（信心 C）。"
                             "官方開放資料無此類設施或有名冊而無座標"),
    ("覆蓋不足", "F9CB9C", "嫌惡設施之部分子欄無資料來源，真實最近設施可能更近 → 推算等級"
                          "僅為樂觀上限，不得據以判級，更不得據以填「無」"),
]

# ---------------- 表3 設施欄位：(名稱格, 核取格, 距離格) ----------------------
_L = lambda r: (f"F{r}", f"I{r}", f"J{r}")
_LH = lambda r: (f"F{r}", f"H{r}", f"J{r}")
_R = lambda r: (f"S{r}", f"T{r}", f"U{r}")
FAC_CELLS = {
    ("near_station", "H13"): _LH(13), ("near_station", "H14"): _LH(14),
    ("near_station", "H15"): _LH(15), ("near_station", "H16"): _LH(16),
    ("near_busstop", "H17"): _LH(17),
    ("near_interchange", "I19"): _L(19),
    ("near_school", "I35"): _L(35), ("near_school", "I36"): _L(36),
    ("near_school", "I37"): _L(37), ("near_school", "I38"): _L(38),
    ("near_market", "I39"): _L(39), ("near_market", "I40"): _L(40),
    ("near_market", "I41"): _L(41),
    ("near_park", "I42"): _L(42), ("near_park", "I43"): _L(43), ("near_park", "I44"): _L(44),
    ("near_tourism", "T4"): _R(4),
    ("parking", "T6"): _R(6),
    ("near_service", "T8a"): _R(8),      # 表3 服務性設施僅一列 → 寫入整體最近者
    ("utility_facility", "T14"): _R(14), ("utility_facility", "T16"): _R(16),
    ("funeral_facility", "T18"): _R(18),
    ("waste_facility", "T22"): _R(22), ("waste_facility", "T23"): _R(23),
    ("waste_facility", "T24"): _R(24),
    ("pollution", "T25"): _R(25), ("pollution", "T26"): _R(26),
    ("pollution", "T27"): _R(27), ("pollution", "T28"): _R(28),
    ("pollution", "T29"): _R(29),
}
RULE = ("手冊伍、一(二)6：距離自本區段中心點起算；"
        "伍、一(二)8(2)：同一細項多設施時取對地價影響最大者（＝最近者）")


def kind_of(conf, usable):
    if not usable:
        return "覆蓋不足"
    return "外部推算" if conf in ("A", "B") else "外部推算OSM"


def tick(text, in_segment):
    if not isinstance(text, str):
        return text
    if in_segment:
        return text.replace("○本區段內", "●本區段內", 1)
    i = text.find("○本區段外")
    return text[:i] + "●本區段外" + text[i + len("○本區段外"):] if i >= 0 else text


def fill_facilities(ws, sn, inf):
    """寫入表3 設施欄位（名稱／內外核取／距離），回傳可用之判級結果"""
    levels = {}
    sc = inf["segments"][sn]
    geo = f"區段中心 TWD97 {sc['center_twd97']} ±{sc['uncertainty_m']}m（由 OSM 界街交會點求得）"
    for code, e in inf["items"].items():
        seg = e["by_segment"].get(sn, {})
        usable = e["coverage"]["usable_for_grading"]
        subs = {s["cell"]: s for s in seg.get("subs", [])}
        for s in e["subs"]:
            key = (code, s["cell"])
            if key not in FAC_CELLS:
                continue
            nc, kc, dc = FAC_CELLS[key]
            if code == "near_service":
                agg = seg.get("aggregated")
                if not agg:
                    continue
                name, dist, conf, src, sub_label = (agg["facility"], agg["distance_m"],
                                                    agg["confidence"], agg["source"], agg["sub"])
            else:
                r = subs.get(s["cell"], {})
                if not r.get("found"):
                    prn = X.T3_SUB_LABELS.get(nc)
                    if prn:
                        ws[nc] = prn
                    X.shade(ws, nc, "資料不足")
                    X.shade(ws, dc, "資料不足")
                    X.basis_rows.append(
                        [sn, nc, f"{e['label']}－{s['name']}", "（留空）",
                         "開放資料無此類設施之可定位記錄",
                         f"{s['note'] or '無對應開放資料'}｜查無資料 ≠ 無設施，"
                         f"不得填「無」（正向設施之「無」＝最劣級，嫌惡設施之「無」＝最優級）",
                         f"{REGION}.regional.{code}", "資料不足"])
                    continue
                name, dist, conf, src, sub_label = (r["nearest"], r["distance_m"],
                                                    r["confidence"], r["source"], r["name"])
            kind = kind_of(conf, usable)
            if name in ("(未命名)", "（未命名）"):
                name = f"{s['name']}（OSM 未命名圖徵）"
            prn = X.T3_SUB_LABELS.get(nc)          # 表單既印之子欄標籤（國小／傳統市場…）
            disp = f"{prn}：{name}" if prn else name
            X.put(ws, nc, disp, kind, f"{e['label']}－{sub_label}",
                  f"{src}（信心 {conf}）", f"{RULE}｜{geo}",
                  f"{REGION}.regional.{code}", seg=sn)
            ws[kc].value = tick(ws[kc].value, dist == 0)
            X.shade(ws, kc, kind)
            c = X.put(ws, dc, dist, kind, f"{e['label']}－{sub_label}（距離 M）",
                      f"{src}（信心 {conf}）",
                      f"區段中心點至最近者之直線距離 {dist} M｜{geo}",
                      f"{REGION}.regional.{code}", seg=sn)
            c.number_format = "0"
            c.alignment = Alignment(horizontal="center")
        if seg.get("rank") and usable:
            levels[code] = {
                "rank": seg["rank"], "label": seg["level_label"],
                "criterion": seg["criterion"],
                "_conf": (seg.get("aggregated") or {}).get("confidence", "C"),
                "_borderline": seg.get("borderline", False),
                "_gap": seg.get("nearest_boundary_gap_m"),
                "_agg": seg.get("aggregated"),
            }
        elif seg.get("rank") and not usable:
            X.basis_rows.append(
                [sn, X.T3_LEVEL_CELL[code][0], e["label"],
                 f"（留空；推算為第{seg['rank']}級 {seg['level_label']}，僅供參考）",
                 f"推算最近者：{(seg.get('aggregated') or {}).get('facility')} "
                 f"{(seg.get('aggregated') or {}).get('distance_m')}M",
                 e["coverage"]["reason"], f"{REGION}.regional.{code}", "覆蓋不足"])
    return levels


def recolor_levels(ws, sn, fac_levels, crit, inf):
    """把設施類的優劣等級格由『AI判定』改標為對應之推算信心底色，並改寫註解"""
    for code, g in fac_levels.items():
        if code not in X.T3_LEVEL_CELL:
            continue
        rank_c, cnt_c = X.T3_LEVEL_CELL[code]
        e = inf["items"][code]
        agg = g.get("_agg") or {}
        kind = kind_of(g["_conf"], True)
        note = (f"【{kind}】{e['label']}（第N級）\n"
                f"填寫值：{g['rank']}（{g['label']}）\n"
                f"最近設施：{agg.get('sub')}／{agg.get('facility')} {agg.get('distance_m')} M"
                f"（{agg.get('source')}，信心 {agg.get('confidence')}）\n"
                f"判定依據：基準表級距「{g['criterion']}」→ 第{g['rank']}／共"
                f"{e['level_count']}級\n"
                f"距最近級距邊界 {g['_gap']} M"
                + ("　⚠️臨界：邊界距離小於區段中心不確定半徑，等級可能翻轉，須人工複核"
                   if g["_borderline"] else "")
                + f"\n最大修正率 ±{e['max_adjustment']}%　級距 {e['step']}%"
                + f"\n基準表項目ID：{REGION}.regional.{code}")
        for cc in (rank_c, cnt_c):
            X.shade(ws, cc, kind)
        cm = Comment(note, X.AUTHOR)
        cm.width, cm.height = 380, 200
        ws[rank_c].comment = cm


def build_table3_v3(ds, segs, seg_order, levels, crit, inf):
    wb = openpyxl.load_workbook(os.path.join(X.IN, "表3地價區段勘查表.xlsx"))
    tpl = wb.active
    tpl.title = seg_order[0]
    sheets = [tpl] + [wb.copy_worksheet(tpl) for _ in seg_order[1:]]
    for ws, sn in zip(sheets, seg_order):
        ws.title = sn
        # 先鋪原表版面與既有 15 細項（含表單既印之設施子欄標籤）
        fac_lv = {c: {"rank": e["by_segment"][sn]["rank"],
                      "label": e["by_segment"][sn]["level_label"],
                      "criterion": e["by_segment"][sn]["criterion"]}
                  for c, e in inf["items"].items()
                  if e["coverage"]["usable_for_grading"]
                  and e["by_segment"].get(sn, {}).get("rank")}
        merged = dict(levels[sn][0])
        merged.update(fac_lv)
        X.fill_table3(ws, ds, segs[sn], sn, merged, levels[sn][1], crit)
        # 再寫入設施名稱／內外／距離（覆蓋在子欄標籤之上，格式「標籤：名稱」）
        fac = fill_facilities(ws, sn, inf)
        recolor_levels(ws, sn, fac, crit, inf)
    X.add_basis_sheet(wb, "表3 地價區段勘查表（新北市樹林區 普通住宅用地，4 個地價區段）"
                          "／設施類欄位為開放資料推算值，非現場勘查")
    p = os.path.join(OUT, "表3_地價區段勘查表_已填.xlsx")
    wb.save(p)
    return p


def main():
    inf = json.load(open(INF, encoding="utf-8"))
    ds = Dataset()
    crit = ds.criteria(REGION, "regional")
    crit_ind = ds.criteria(REGION, "individual")
    segs = ds.segments(REGION)
    case = list(ds.cases(REGION).values())[0]
    seg_order = case["table5"]["segment_nos"]
    levels = {sn: derive_segment_levels(ds, REGION, segs[sn]) for sn in seg_order}
    t4i = derive_table4_individual(ds, REGION, case, segs)

    # 合併：既有 15 細項 ＋ 開放資料推算之設施類細項
    fac_by_seg = {}
    for sn in seg_order:
        fac_by_seg[sn] = {c: {"rank": e["by_segment"][sn]["rank"],
                              "label": e["by_segment"][sn]["level_label"],
                              "criterion": e["by_segment"][sn]["criterion"]}
                          for c, e in inf["items"].items()
                          if e["coverage"]["usable_for_grading"]
                          and e["by_segment"].get(sn, {}).get("rank")}
    merged_levels = {sn: ({**levels[sn][0], **fac_by_seg[sn]}, levels[sn][1])
                     for sn in seg_order}

    X.basis_rows = []
    p3 = build_table3_v3(ds, segs, seg_order, levels, crit, inf)
    n3 = len(X.basis_rows)

    X.basis_rows = []
    p5, group_ok, group_vals = X.build_table5(ds, case, segs, seg_order, merged_levels, crit)
    n5 = len(X.basis_rows)

    X.basis_rows = []
    p4 = X.build_table4(ds, case, segs, seg_order, merged_levels, t4i, crit_ind)
    n4 = len(X.basis_rows)

    print(f"表3 → {os.path.relpath(p3, ROOT)}（{len(seg_order)} 個區段工作表，{n3} 筆依據）")
    print(f"表4 → {os.path.relpath(p4, ROOT)}（{n4} 筆依據）")
    print(f"表5 → {os.path.relpath(p5, ROOT)}（{n5} 筆依據）")
    ok = sorted(g for g in group_ok if group_ok[g])
    ng = sorted(g for g in group_ok if not group_ok[g])
    print("表5 小計成立之主要項目：", ok)
    print("表5 小計不成立之主要項目：", ng)
    for g in ok:
        print(f"   ({g}) 小計 =", [round(v, 2) for v in group_vals[g]] if group_vals.get(g) else "-")
    return {"table3": p3, "table4": p4, "table5": p5,
            "group_ok": group_ok, "group_vals": group_vals,
            "levels": merged_levels, "inference": inf}


if __name__ == "__main__":
    main()
